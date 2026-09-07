# -*- coding: utf-8 -*-
"""Generate Dify-friendly Excel for 人生旅程 20 patterns."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "docs" / "人生旅程-20种模式-知识库-v2.xlsx"
OUT_LEGACY = Path(__file__).resolve().parents[1] / "docs" / "人生旅程-20种模式-知识库.xlsx"

DISCLAIMER = (
    "本模式仅作体验式家庭关系探索的参照假设，不作为心理诊断、人格定型、"
    "婚姻优劣判定、亲子关系定级或财富预测。用户个人联想与真实处境优先于系统词典。"
)

# 产品口径（与 C 端 v1.0.4 一致）：猴子 = 参与者本人，不是重要他人
SYMBOLS = (
    "猴子=参与者本人（「我/自己」的自我象征，同指）；"
    "小鸟=孩子或需要照顾的人；背包=责任/义务/负担；"
    "蛇=财富/资源/安全感（也可能含风险）。"
    "用户说「猴子」不必再说「我」；匹配时猴子与我/自己视为同一主体。"
)

ROLE_ALIGN = (
    "归一规则：用户文本中的「猴子」「我」「自己」均指参与者本人。"
    "对齐典型安排前先做同指归一：用户说「猴子背背包」=「自己背背包」。"
    "知识库典型安排里若同时出现「猴子」与「自己」描述参与者动作，按同一主体理解，"
    "不要拆成两个人。模式名含「重要他人」时，看的是位置结构（跟随/不在最前/责任外置），"
    "不是把猴子当成另一位他人。"
)

ROWS = [
    {
        "id": "j01",
        "name": "全部自己承载型",
        "arrange": "自己背背包、抱着小鸟、同时管理蛇，四件都由自己携带或牵管（用户可说「猴子」代指自己）。",
        "keywords": "自己背背包,猴子背背包,抱着小鸟,同时管理蛇,一个人全扛,四件都在身上",
        "confuse": "责任过载型(j02)：也是自己背包，但未强调四件同时由自己携带/管理。家庭维系型(j08)：强调在中间协调牵带，未必四件全在身上。孤立承担型(j19)：自己独自背包，其他对象距离较远。",
        "positive": "责任感较强；重视家庭与关系；愿意保护和照顾他人。",
        "watch": "是否习惯把所有事情都揽在自己身上；是否不容易信任别人或向别人求助；是否长期忽略了自己的疲惫和需要。",
    },
    {
        "id": "j02",
        "name": "责任过载型",
        "arrange": "背包很重，由自己背着，同时还要照顾其他对象。",
        "keywords": "背包很重,自己背着,猴子背着包,还要照顾其他对象,负担过重,责任太大",
        "confuse": "全部自己承载型(j01)：强调四件同时由自己携带/管理。责任中心型(j13)：背包在队伍中心、众人围着任务转。孤立承担型(j19)：自己背包但其他人离得远，未必同时照顾对象。",
        "positive": "做事可靠；具有责任意识；面对困难时不轻易放弃。",
        "watch": "是否承担了超出自己能力或角色范围的责任；是否把“负责”等同于“必须一个人完成”；是否已经出现疲惫、委屈或无力感。",
    },
    {
        "id": "j03",
        "name": "孩子中心型",
        "arrange": "小鸟位于最前面、中间或离自己最近，其他对象围绕小鸟行动。",
        "keywords": "小鸟最前面,小鸟在中间,小鸟离自己最近,其他对象围绕小鸟,孩子在中心",
        "confuse": "亲子同行型(j07)：自己与小鸟并肩，小鸟可独立行动，不是全员围着小鸟。孩子承责型(j11)：小鸟背背包/照顾他人，是孩子在承责。亲子联盟型(j09)：小鸟与其他象征成一组，自己稍在侧后。",
        "positive": "高度重视孩子；对孩子的需要比较敏感；愿意为孩子提供保护和支持。",
        "watch": "家庭生活是否过度围绕孩子展开；成人是否逐渐失去自己的生活方向；孩子是否被放在了不适合其年龄的中心位置。",
    },
    {
        "id": "j04",
        "name": "重要他人引领型",
        "arrange": "自己不在队伍最前，落在后方跟随；小鸟、蛇或路线方向相对更靠前（跟随姿态，非自我主导）。",
        "keywords": "跟在后面,落在后面,不在最前面,跟随,他人在前,不走在最前,引路在前方",
        "confuse": "自我主导型(j05)：自己/猴子走在最前面。自我后置型(j18)：自己在最后但是照看全队，不是单纯跟随。灵活调整型(j20)：位置随情况变化。",
        "positive": "愿意信任重要他人；能够接受他人的带领；重视合作与分工。",
        "watch": "自己是主动信任还是被动跟随；是否过度依赖对方决定方向；自己是否能够表达不同意见。",
    },
    {
        "id": "j05",
        "name": "自我主导型",
        "arrange": "自己（用户可说猴子）走在最前面，其他对象跟随。",
        "keywords": "自己走在最前面,猴子走在最前面,其他人跟随,自己带路,猴子带路,走在最前",
        "confuse": "重要他人引领型(j04)：自己不在最前、在后方跟随。自我后置型(j18)：自己走在最后照看所有人。灵活调整型(j20)：位置随路况变化。",
        "positive": "有方向感和决策能力；愿意承担家庭或领导责任；面对复杂情况时比较主动。",
        "watch": "是否习惯替所有人作决定；是否允许他人表达不同意见；是否因不信任别人而承担过多。",
    },
    {
        "id": "j06",
        "name": "平等伙伴型",
        "arrange": "自己与小鸟或其他旅伴并肩同行，共同分担责任，不是一人独揽。",
        "keywords": "并肩,共同分担,一起走,平等合作,与小鸟并肩,与蛇并肩",
        "confuse": "自我主导型(j05)：自己/猴子单独在最前领队。家庭维系型(j08)：自己在中间分别协调两边，不是并肩分担。",
        "positive": "重视平等、合作与协商；愿意与重要他人共同承担。",
        "watch": "表面并肩是否真正意味着平等；谁拥有最终决定权；双方的责任分配是否公平。",
    },
    {
        "id": "j07",
        "name": "亲子同行型",
        "arrange": "自己与小鸟并肩，小鸟可以独立行动。",
        "keywords": "自己与小鸟并肩,小鸟可独立行动,陪伴不替代,亲子并肩",
        "confuse": "孩子中心型(j03)：全员围绕小鸟，小鸟在最前/中间。家庭维系型(j08)：自己在中间协调多对象。亲子联盟型(j09)：小鸟与其他象征成组、自己靠边。",
        "positive": "尊重孩子的主体性；重视陪伴而非替代；愿意给孩子成长的空间。",
        "watch": "是否真正允许孩子有自己的方向；孩子遇到困难的时候，自己会陪伴还是立即接管；并肩是否符合孩子的年龄与发展需要。",
    },
    {
        "id": "j08",
        "name": "家庭维系型",
        "arrange": "自己处于中间，同时带着小鸟和蛇（或一手牵鸟、一手管蛇/背包），协调两边。",
        "keywords": "自己在中间,一手牵小鸟,协调,和事佬,同时带着,中间位置",
        "confuse": "全部自己承载型(j01)：四件同时由自己携带/管理，负荷更全。平等伙伴型(j06)：并肩分担而非居中协调。高度融合型(j15)：全部紧贴不许分开。",
        "positive": "重视家庭完整与情感连接；愿意协调家庭成员之间的关系；具有较强的家庭责任感。",
        "watch": "是否把维系家庭关系视为自己一个人的责任；是否长期处于协调者或“和事佬”的位置；自己是否拥有独立的空间。",
    },
    {
        "id": "j09",
        "name": "亲子联盟型",
        "arrange": "小鸟与其他象征（如蛇/背包）形成一组，自己稍在侧旁或后方。",
        "keywords": "小鸟与其他成一组,自己在旁边,自己在后面,被排除在外,稍靠后",
        "confuse": "亲子同行型(j07)：自己与小鸟并肩。孩子中心型(j03)：小鸟被全员围绕。家庭维系型(j08)：自己在中间协调。",
        "positive": "重要他人与孩子关系较亲近；家庭成员之间能够相互支持；自己愿意让出部分的照顾空间。",
        "watch": "自己是否感到被排除；家庭中是否形成固定的联盟；成人关系是否被亲子关系取代。",
    },
    {
        "id": "j10",
        "name": "伴侣承责型",
        "arrange": "背包主要不在自己身上，或责任似乎由他人/外部承担；自己较少背背包。",
        "keywords": "背包不在自己身上,背包在脚边,不是自己背,责任交给他人,期待别人承担,较少自己背",
        "confuse": "自我主导/责任过载(j02/j05)：自己/猴子明确背背包。全部自己承载(j01)：四件都在自己身上。孤立承担(j19)：自己独自背包、他人远离。",
        "positive": "愿意让重要他人分担；对对方具有信任和责任期待。",
        "watch": "对方是否愿意承担；责任分配是否公平；自己是否把过多的责任交给对方。",
    },
    {
        "id": "j11",
        "name": "孩子承责型",
        "arrange": "小鸟背着背包、照顾自己（或管理蛇/背包）。",
        "keywords": "小鸟背背包,小鸟照顾自己,小鸟管理蛇,孩子承担责任,角色倒置,亲职化",
        "confuse": "孩子中心型(j03)：小鸟被围绕、被保护。责任过载型(j02)：过载的是自己。伴侣承责型(j10)：背包不在自己身上。",
        "positive": "孩子可能被赋予了超出年龄的责任；孩子可能需要照顾成人的情绪；家庭中可能存在着角色倒置或亲职化的倾向。",
        "watch": "孩子是否经常安慰父母；是否承担协调家庭矛盾的任务；是否因为“懂事”而压抑自己的需要。",
    },
    {
        "id": "j12",
        "name": "财富中心型",
        "arrange": "蛇走在前面、位于中间、紧贴自己或受到重点保护。",
        "keywords": "蛇走在前面,蛇在中间,蛇紧贴自己,重点保护蛇,财富在中心",
        "confuse": "责任中心型(j13)：队伍中心是背包不是蛇。控制保护型(j14)：多个对象被牵着/拴着/抱着。",
        "positive": "重视经济安全；对资源与机会比较敏感；具有财富管理的意识。",
        "watch": "财富是否过度影响关系和人生选择；是否容易因经济问题焦虑；是否将经济能力与个人价值联系起来。",
    },
    {
        "id": "j13",
        "name": "责任中心型",
        "arrange": "背包位于队伍中心，所有人的行动都围绕着背包展开。",
        "keywords": "背包在队伍中心,围绕背包行动,任务驱动,应该做什么",
        "confuse": "责任过载型(j02)：背包很重且由自己背。财富中心型(j12)：中心物是蛇。全部自己承载型(j01)：自己同时携带/manage 多件。",
        "positive": "重视责任与任务；做事有计划；不轻易逃避义务。",
        "watch": "生活是否长期被“应该做什么”推动；是否忽略关系、感受和休息；是否将个人价值等同于完成任务。",
    },
    {
        "id": "j14",
        "name": "控制保护型",
        "arrange": "多个对象被牵着、拴着、抱着或装进容器。",
        "keywords": "牵着,拴着,抱着,装进容器,不让离开视线,保护变成限制",
        "confuse": "高度融合型(j15)：紧贴、牵连、不允许分开。全部自己承载型(j01)：自己携带多件，未必用拴/装的方式。",
        "positive": "安全意识较强；希望保护重要的人和资源；对风险具有较高敏感度。",
        "watch": "保护是否逐渐变成限制；是否难以接受他人离开自己的视线；是否因为担心失控而减少他人的自主空间。",
    },
    {
        "id": "j15",
        "name": "高度融合型",
        "arrange": "所有对象都紧贴自己、彼此牵连，不允许分开。",
        "keywords": "全部紧贴自己,彼此牵连,不允许分开,没有距离,粘在一起",
        "confuse": "控制保护型(j14)：用牵/拴/抱/装容器来限制。边界清晰型(j16)：保持联系但能独立行动。",
        "positive": "重视亲密和家庭完整；希望大家相互陪伴；具有较强的归属感。",
        "watch": "是否对分离比较敏感；是否难以允许家庭成员独立；个人边界是否容易模糊。",
    },
    {
        "id": "j16",
        "name": "边界清晰型",
        "arrange": "各对象保持联系，但能够独立行动，责任分配相对清楚。",
        "keywords": "保持联系又能独立,责任分配清楚,有连接有边界,各走各的但仍相关",
        "confuse": "各自独立型(j17)：彼此距离较远、各自行动，连接更弱。高度融合型(j15)：不允许分开。",
        "positive": "能够兼顾亲密与独立；尊重不同成员的空间；具有较清晰的责任边界。",
        "watch": "独立是否伴随足够的情感连接；困难出现时是否能够相互求助；边界是否经过共同协商。",
    },
    {
        "id": "j17",
        "name": "各自独立型",
        "arrange": "小鸟、蛇等各走各的，彼此距离较远（自己/猴子未紧挨着它们）。",
        "keywords": "小鸟自己飞,蛇自己爬,彼此距离较远,各自行动,不挨着",
        "confuse": "边界清晰型(j16)：仍保持联系且责任清楚。孤立承担型(j19)：自己独自背包、其他人远。灵活调整型(j20)：会随情况改位置。",
        "positive": "相信每个人具有独立能力；尊重自主选择；不愿过度控制他人。",
        "watch": "独立是否变成了情感疏离；家庭成员之间是否缺少支持。",
    },
    {
        "id": "j18",
        "name": "自我后置型",
        "arrange": "自己走在最后，负责照看所有对象。",
        "keywords": "自己走在最后,猴子走在最后,在后面照看,收尾,把别人放前面",
        "confuse": "自我主导型(j05)：自己/猴子走最前。重要他人引领型(j04)：自己在后但是跟随，未必是照看全队。孤立承担型(j19)：其他人距离远。",
        "positive": "重视他人；能够照顾全局；具有支持与收尾能力。",
        "watch": "是否习惯把自己的需要放在最后；是否总担心别人掉队；是否长期处于疲惫、孤独或被忽视的状态。",
    },
    {
        "id": "j19",
        "name": "孤立承担型",
        "arrange": "自己独自背着背包，其他对象距离较远。",
        "keywords": "自己独自背背包,猴子独自背背包,其他对象距离较远,一个人扛,身边没人",
        "confuse": "全部自己承载型(j01)：其他对象仍在自己身上（抱/牵/管）。责任过载型(j02)：还要照顾其他对象。各自独立型(j17)：大家各自走。",
        "positive": "独立能力较强；面对压力时能够坚持；不轻易依赖别人。",
        "watch": "是否不容易信任或求助；是否习惯强撑；是否在关系中感到孤独。",
    },
    {
        "id": "j20",
        "name": "灵活调整型",
        "arrange": "根据道路、天气和不同情况不断调整位置与分工。",
        "keywords": "根据路况调整,根据天气调整,不断换位置,动态分工,协商调整",
        "confuse": "仅当描述明确出现“随情况改变位置/分工”时才选此项。提到山、桥、阳光等风景不够。固定谁走最前/谁背包应优先匹配 j01–j19。",
        "positive": "适应能力较强；能够根据现实情况调整策略；重视协商与动态平衡。",
        "watch": "调整是否建立在清晰原则之上；是否为了照顾别人而不断改变自己；谁在决定每次调整。",
    },
]


def build_content(row):
    return "\n".join(
        [
            "测评类型: journey",
            f"patternId: {row['id']}",
            f"模式名称: {row['name']}",
            f"象征词典: {SYMBOLS}",
            f"角色对齐: {ROLE_ALIGN}",
            f"匹配关键词: {row['keywords']}",
            f"典型安排: {row['arrange']}",
            f"不要与这些混淆: {row['confuse']}",
            f"资源取向: {row['positive']}",
            f"可留意: {row['watch']}",
            f"使用说明: {DISCLAIMER}",
        ]
    )


def style_header(ws, ncols):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF", name="Microsoft YaHei", size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def style_body(ws, nrows, ncols, wrap_cols):
    thin = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )
    font = Font(name="Microsoft YaHei", size=10)
    wrap = Alignment(vertical="top", wrap_text=True)
    top = Alignment(vertical="top", wrap_text=False)
    alt = PatternFill("solid", fgColor="F6F8FA")
    for r in range(2, nrows + 1):
        ws.row_dimensions[r].height = 88 if ws.title == "知识库正文" else 64
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.font = font
            cell.border = thin
            cell.alignment = wrap if c in wrap_cols else top
            if r % 2 == 0:
                cell.fill = alt


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "知识库正文"
    headers1 = [
        "content",
        "assessmentType",
        "patternId",
        "模式名称",
        "典型安排",
        "匹配关键词",
        "不要与这些混淆",
        "资源取向",
        "可留意",
        "使用说明",
    ]
    ws1.append(headers1)
    for row in ROWS:
        ws1.append(
            [
                build_content(row),
                "journey",
                row["id"],
                row["name"],
                row["arrange"],
                row["keywords"],
                row["confuse"],
                row["positive"],
                row["watch"],
                DISCLAIMER,
            ]
        )
    style_header(ws1, len(headers1))
    style_body(ws1, 1 + len(ROWS), len(headers1), wrap_cols={1, 5, 6, 7, 8, 9, 10})
    autosize(ws1, [70, 16, 12, 18, 36, 36, 48, 36, 40, 40])
    ws1.sheet_properties.tabColor = "1F4E79"

    ws2 = wb.create_sheet("模式索引")
    headers2 = ["patternId", "模式名称", "一句话匹配要点"]
    ws2.append(headers2)
    for row in ROWS:
        ws2.append([row["id"], row["name"], row["arrange"]])
    style_header(ws2, 3)
    style_body(ws2, 1 + len(ROWS), 3, wrap_cols={3})
    autosize(ws2, [14, 20, 56])
    for r in range(2, 22):
        ws2.row_dimensions[r].height = 32
    ws2.sheet_properties.tabColor = "2E7D32"

    ws3 = wb.create_sheet("Dify导入说明")
    ws3["A1"] = "人生旅程 20 种模式 · Dify 知识库导入说明"
    ws3["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color="1F4E79")
    guide = [
        "",
        "一、这份表怎么用",
        "1. 打开 Dify → 知识库 → 导入 → 选择本 Excel。",
        "2. 分段方式选「按行」。每一行 = 一种模式，共 20 段。",
        "3. 正文列请选「content」（第 1 列）。",
        "4. 若支持元数据：assessmentType、patternId、模式名称。",
        "",
        "二、象征词典（产品口径，全局一致）",
        SYMBOLS,
        ROLE_ALIGN,
        "",
        "三、与旧版业务文档的差异",
        "旧版活动文档曾写「猴子=重要他人」。本测评 C 端与 Dify 知识库统一为：猴子=我/自己。",
        "导入后请同步更新匹配 LLM 与生图提示词。",
        "",
        "四、数据来源",
        "journey-patterns.js + v1.0.4 产品口径；典型安排/关键词已按猴子=自我象征重写。",
    ]
    for i, line in enumerate(guide, 2):
        ws3.cell(i, 1, line)
        ws3.cell(i, 1).font = Font(name="Microsoft YaHei", size=11)
        ws3.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws3.row_dimensions[i].height = 22 if line else 10
    ws3.column_dimensions["A"].width = 110
    ws3.sheet_properties.tabColor = "F57C00"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    try:
        wb.save(OUT_LEGACY)
        print(f"wrote {OUT_LEGACY} rows={len(ROWS)}")
    except PermissionError:
        print(f"wrote {OUT} rows={len(ROWS)} (legacy locked, close Excel to overwrite)")
        return
    print(f"wrote {OUT} rows={len(ROWS)}")


if __name__ == "__main__":
    main()
